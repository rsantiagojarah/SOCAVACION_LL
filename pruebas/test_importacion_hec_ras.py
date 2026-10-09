"""Pruebas TDD de importación de secciones HEC-RAS desde XLSX."""
from pathlib import Path
import os
import sys
import tempfile
import unittest

from openpyxl import Workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))


def crear_libro(ruta: Path, *, desplazamiento=0, aguas=("WS PF 1",)):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Sección 101"
    columna = 1 + desplazamiento
    hoja.cell(2, columna + 1, "Ground")
    hoja.cell(3, columna, "x")
    hoja.cell(3, columna + 1, "y")
    for fila, punto in enumerate(((0.0, 10.0), (5.0, 7.0), (12.0, 11.0)), 4):
        hoja.cell(fila, columna, punto[0])
        hoja.cell(fila, columna + 1, punto[1])
    columna += 2
    hoja.cell(2, columna + 1, "Bank Sta")
    hoja.cell(3, columna, "x")
    hoja.cell(3, columna + 1, "y")
    hoja.cell(4, columna, 1.25)
    hoja.cell(4, columna + 1, 9.25)
    hoja.cell(5, columna, 10.5)
    hoja.cell(5, columna + 1, 10.1)
    for indice, nombre in enumerate(aguas):
        columna += 2
        hoja.cell(2, columna + 1, nombre)
        hoja.cell(3, columna, "x")
        hoja.cell(3, columna + 1, "y")
        cota = 9.0 + indice
        hoja.cell(4, columna, 0.0)
        hoja.cell(4, columna + 1, cota)
        hoja.cell(6, columna, 12.0)
        hoja.cell(6, columna + 1, cota)
    libro.save(ruta)
    libro.close()


class PruebasLectorHecRas(unittest.TestCase):
    def setUp(self):
        try:
            from socavacion_ll.infraestructura.hec_ras_xlsx import leer_seccion_hec_ras
        except ModuleNotFoundError:
            self.fail("Falta el lector XLSX de secciones HEC-RAS")
        self.leer = leer_seccion_hec_ras

    def test_detecta_series_sin_depender_de_columnas_fijas(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "seccion.xlsx"
            crear_libro(ruta, desplazamiento=3, aguas=("WS PF 1", "WS PF 2"))
            seccion = self.leer(ruta)
        self.assertEqual(seccion.hoja, "Sección 101")
        self.assertEqual([(p.abscisa_m, p.cota_m) for p in seccion.puntos_terreno],
                         [(0.0, 10.0), (5.0, 7.0), (12.0, 11.0)])
        self.assertEqual([(p.nombre, p.cota_m) for p in seccion.perfiles_agua],
                         [("WS PF 1", 9.0), ("WS PF 2", 10.0)])
        self.assertEqual(seccion.estaciones_margen_m, (1.25, 10.5))

    def test_rechaza_extension_y_mas_de_una_hoja_con_ground(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "seccion.xlsx"
            crear_libro(ruta)
            libro = Workbook()
            for hoja in (libro.active, libro.create_sheet("Otra")):
                hoja.cell(2, 2, "Ground")
                hoja.cell(3, 1, "x")
                hoja.cell(3, 2, "y")
            doble = Path(directorio) / "doble.xlsx"
            libro.save(doble)
            libro.close()
            with self.assertRaisesRegex(ValueError, "una sola hoja"):
                self.leer(doble)
            with self.assertRaisesRegex(ValueError, "XLSX"):
                self.leer(ruta.with_suffix(".xls"))

    def test_rechaza_pares_incompletos_formulas_y_abscisas_no_crecientes(self):
        casos = ((None, "par incompleto"), ("=1+1", "fórmulas"), (-1.0, "decrecientes"))
        for valor, mensaje in casos:
            with self.subTest(valor=valor), tempfile.TemporaryDirectory() as directorio:
                ruta = Path(directorio) / "invalida.xlsx"
                crear_libro(ruta)
                from openpyxl import load_workbook
                libro = load_workbook(ruta)
                hoja = libro.active
                if valor is None:
                    hoja.cell(5, 2).value = None
                elif isinstance(valor, str):
                    hoja.cell(5, 1).value = valor
                else:
                    hoja.cell(5, 1).value = valor
                libro.save(ruta)
                libro.close()
                with self.assertRaisesRegex(ValueError, mensaje):
                    self.leer(ruta)

    def test_conserva_un_tramo_vertical_sin_desplazarlo(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "vertical.xlsx"
            crear_libro(ruta)
            from openpyxl import load_workbook
            libro = load_workbook(ruta)
            hoja = libro.active
            hoja.cell(5, 1).value = 0.0
            libro.save(ruta)
            libro.close()
            seccion = self.leer(ruta)
        self.assertEqual(seccion.puntos_terreno[0].abscisa_m, 0.0)
        self.assertEqual(seccion.puntos_terreno[1].abscisa_m, 0.0)

    def test_rechaza_perfil_de_agua_no_horizontal(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "inclinada.xlsx"
            crear_libro(ruta)
            from openpyxl import load_workbook
            libro = load_workbook(ruta)
            hoja = libro.active
            hoja.cell(6, 6).value = 9.0001
            libro.save(ruta)
            libro.close()
            with self.assertRaisesRegex(ValueError, "horizontal"):
                self.leer(ruta)

    def test_rechaza_mas_de_doscientos_puntos(self):
        with tempfile.TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "grande.xlsx"
            crear_libro(ruta)
            from openpyxl import load_workbook
            libro = load_workbook(ruta)
            hoja = libro.active
            for fila in range(4, 205):
                hoja.cell(fila, 1, float(fila))
                hoja.cell(fila, 2, 7.0)
            libro.save(ruta)
            libro.close()
            with self.assertRaisesRegex(ValueError, "200"):
                self.leer(ruta)


class PruebasInterfazImportacion(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from PySide6.QtWidgets import QApplication
        cls.aplicacion = QApplication.instance() or QApplication([])

    def test_exige_confirmar_metros_y_permite_elegir_perfil(self):
        from socavacion_ll.dominio.modelos import PuntoTerreno
        from socavacion_ll.infraestructura.hec_ras_xlsx import (
            PerfilAguaHecRas, SeccionHecRas,
        )
        try:
            from socavacion_ll.presentacion.importacion_hec_ras import DialogoImportacionHecRas
        except ModuleNotFoundError:
            self.fail("Falta el diálogo de confirmación de la importación HEC-RAS")
        seccion = SeccionHecRas(
            "seccion.xlsx", "RS 101",
            (PuntoTerreno(0, 10), PuntoTerreno(12, 11)),
            (PerfilAguaHecRas("WS PF 1", 9), PerfilAguaHecRas("WS PF 2", 10)),
            (1.25, 10.5),
        )
        dialogo = DialogoImportacionHecRas(seccion)
        self.assertFalse(dialogo.boton_aceptar.isEnabled())
        self.assertIn("2 puntos", dialogo.resumen.text())
        dialogo.perfiles.setCurrentIndex(1)
        self.assertEqual(dialogo.perfil_seleccionado().cota_m, 10)
        dialogo.confirmar_metros.setChecked(True)
        self.assertTrue(dialogo.boton_aceptar.isEnabled())
        self.assertIn("WS PF 2", dialogo.nota_procedencia())
        self.assertIn("1.25; 10.5", dialogo.nota_procedencia())
        dialogo.deleteLater()

    def test_aplicacion_atomica_preserva_los_demas_datos(self):
        from dataclasses import replace
        from socavacion_ll.dominio.modelos import PuntoTerreno, proyecto_ejemplo
        from socavacion_ll.presentacion.ingreso import PaginaIngreso
        pagina = PaginaIngreso()
        original = replace(proyecto_ejemplo(), notas="Dato anterior")
        pagina.cargar(original)
        puntos = (PuntoTerreno(0, 10), PuntoTerreno(12, 11))
        pagina.establecer_importacion_hec_ras(puntos, 9.0, "Dato anterior\n\n[HEC-RAS]")
        resultado = pagina.obtener_proyecto()
        self.assertEqual(resultado.puntos_terreno, puntos)
        self.assertEqual(resultado.cota_agua_m, 9.0)
        self.assertEqual(resultado.notas, "Dato anterior\n\n[HEC-RAS]")
        self.assertEqual(resultado.caudal_m3s, original.caudal_m3s)
        self.assertEqual(resultado.diametro_medio_mm, original.diametro_medio_mm)
        pagina.deleteLater()

    def test_editor_retiene_importacion_hasta_aceptar(self):
        from socavacion_ll.dominio.modelos import PuntoTerreno
        from socavacion_ll.infraestructura.hec_ras_xlsx import (
            PerfilAguaHecRas, SeccionHecRas,
        )
        from socavacion_ll.presentacion.seccion_irregular import DialogoSeccionIrregular
        seccion = SeccionHecRas(
            "seccion.xlsx", "RS 101",
            (PuntoTerreno(0, 10), PuntoTerreno(12, 11)),
            (PerfilAguaHecRas("WS PF 1", 9),), (1.25, 10.5),
        )

        class Confirmacion:
            def perfil_seleccionado(self):
                return seccion.perfiles_agua[0]

            def nota_procedencia(self):
                return "[Importación HEC-RAS XLSX]"

        dialogo = DialogoSeccionIrregular((), "Nota previa")
        self.assertIn("HEC-RAS", dialogo.boton_importar.text())
        dialogo._aplicar_importacion_confirmada(seccion, Confirmacion())
        self.assertEqual(dialogo.obtener_puntos(), seccion.puntos_terreno)
        self.assertEqual(dialogo.obtener_importacion(),
                         (9, "Nota previa\n\n[Importación HEC-RAS XLSX]"))
        dialogo.deleteLater()

    def test_editor_reparte_el_ancho_entre_abscisa_y_cota(self):
        from PySide6.QtWidgets import QHeaderView
        from socavacion_ll.presentacion.seccion_irregular import DialogoSeccionIrregular
        dialogo = DialogoSeccionIrregular(())
        cabecera = dialogo.tabla.horizontalHeader()
        self.assertEqual(cabecera.sectionResizeMode(0), QHeaderView.ResizeMode.Stretch)
        self.assertEqual(cabecera.sectionResizeMode(1), QHeaderView.ResizeMode.Stretch)
        dialogo.deleteLater()


if __name__ == "__main__":
    unittest.main()
