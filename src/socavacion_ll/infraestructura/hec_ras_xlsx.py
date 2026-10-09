"""Lectura segura y sin conversiones de secciones exportadas por HEC-RAS."""
from dataclasses import dataclass
import math
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from socavacion_ll.dominio.modelos import PuntoTerreno

TAMANO_MAXIMO_BYTES = 10 * 1024 * 1024
PUNTOS_MAXIMOS = 200
TOLERANCIA_AGUA_M = 1e-9


@dataclass(frozen=True)
class PerfilAguaHecRas:
    nombre: str
    cota_m: float


@dataclass(frozen=True)
class SeccionHecRas:
    archivo: str
    hoja: str
    puntos_terreno: tuple[PuntoTerreno, ...]
    perfiles_agua: tuple[PerfilAguaHecRas, ...]
    estaciones_margen_m: tuple[float, ...]


def _texto(valor) -> str:
    return " ".join(str(valor).strip().split()) if valor is not None else ""


def _series(hoja) -> dict[str, tuple[int, int]]:
    encontradas = {}
    for columna in range(2, hoja.max_column + 1):
        nombre = _texto(hoja.cell(2, columna).value)
        eje_x = _texto(hoja.cell(3, columna - 1).value).lower()
        eje_y = _texto(hoja.cell(3, columna).value).lower()
        if nombre and eje_x == "x" and eje_y == "y":
            encontradas[nombre] = (columna - 1, columna)
    return encontradas


def _numero(valor, hoja: str, celda: str) -> float:
    if isinstance(valor, str) and valor.lstrip().startswith("="):
        raise ValueError(f"{hoja}!{celda}: no se admiten fórmulas; exporte valores.")
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        raise ValueError(f"{hoja}!{celda}: se esperaba un número.")
    numero = float(valor)
    if not math.isfinite(numero):
        raise ValueError(f"{hoja}!{celda}: el valor debe ser finito.")
    return numero


def _pares(hoja, columnas: tuple[int, int], nombre: str) -> tuple[tuple[float, float], ...]:
    pares = []
    columna_x, columna_y = columnas
    for fila in range(4, hoja.max_row + 1):
        celda_x = hoja.cell(fila, columna_x)
        celda_y = hoja.cell(fila, columna_y)
        if celda_x.value is None and celda_y.value is None:
            continue
        if celda_x.value is None or celda_y.value is None:
            raise ValueError(
                f"{hoja.title}, {nombre}, fila {fila}: par incompleto de x e y."
            )
        pares.append((_numero(celda_x.value, hoja.title, celda_x.coordinate),
                      _numero(celda_y.value, hoja.title, celda_y.coordinate)))
    return tuple(pares)


def _validar_terreno(pares) -> tuple[PuntoTerreno, ...]:
    if not 2 <= len(pares) <= PUNTOS_MAXIMOS:
        raise ValueError("Ground debe contener entre 2 y 200 puntos.")
    if any(actual[0] < anterior[0] for anterior, actual in zip(pares, pares[1:])):
        raise ValueError("Las abscisas de Ground no pueden ser decrecientes.")
    if any(actual == anterior for anterior, actual in zip(pares, pares[1:])):
        raise ValueError("Ground contiene puntos consecutivos idénticos.")
    return tuple(PuntoTerreno(x, y) for x, y in pares)


def _perfiles_agua(hoja, series) -> tuple[PerfilAguaHecRas, ...]:
    perfiles = []
    for nombre, columnas in series.items():
        if not nombre.upper().startswith("WS "):
            continue
        pares = _pares(hoja, columnas, nombre)
        if not pares:
            raise ValueError(f"El perfil {nombre} no contiene puntos.")
        cota = pares[0][1]
        if any(abs(y - cota) > TOLERANCIA_AGUA_M for _, y in pares[1:]):
            raise ValueError(f"El perfil {nombre} no es horizontal.")
        perfiles.append(PerfilAguaHecRas(nombre, cota))
    if not perfiles:
        raise ValueError("No se encontró ningún perfil de agua WS en la hoja Ground.")
    return tuple(perfiles)


def leer_seccion_hec_ras(ruta) -> SeccionHecRas:
    """Lee una sección XLSX; conserva números y orden sin interpolar ni convertir."""
    archivo = Path(ruta)
    if archivo.suffix.lower() != ".xlsx":
        raise ValueError("Seleccione un archivo HEC-RAS con extensión XLSX.")
    try:
        tamano = archivo.stat().st_size
    except OSError as error:
        raise ValueError(f"No se puede acceder al archivo XLSX: {error}") from error
    if not 0 < tamano <= TAMANO_MAXIMO_BYTES:
        raise ValueError("El archivo XLSX debe ser válido y no superar 10 MiB.")
    try:
        libro = load_workbook(archivo, read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ValueError(f"No se pudo abrir el XLSX: {error}") from error
    try:
        candidatas = [(hoja, _series(hoja)) for hoja in libro.worksheets]
        candidatas = [(hoja, series) for hoja, series in candidatas if "Ground" in series]
        if len(candidatas) != 1:
            raise ValueError("El XLSX debe contener una sola hoja con la serie Ground.")
        hoja, series = candidatas[0]
        terreno = _validar_terreno(_pares(hoja, series["Ground"], "Ground"))
        perfiles = _perfiles_agua(hoja, series)
        margenes = ()
        if "Bank Sta" in series:
            margenes = tuple(x for x, _ in _pares(hoja, series["Bank Sta"], "Bank Sta"))
        return SeccionHecRas(archivo.name, hoja.title, terreno, perfiles, margenes)
    finally:
        libro.close()
