"""Lectura de abertura y porcentaje retenido parcial. No calcula Dm."""
from dataclasses import dataclass
import math
from pathlib import Path
import unicodedata
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

TAMANO_MAXIMO_BYTES = 10 * 1024 * 1024


@dataclass(frozen=True)
class LecturaAberturas:
    archivo: str
    tamices: tuple[tuple[float, float], ...]
    fondo: float


def leer_aberturas_retenido(ruta) -> LecturaAberturas:
    """Leer la primera hoja: abertura (mm), % retenido parcial y, si existe, Fondo."""
    origen = Path(ruta)
    if not origen.is_file():
        raise ValueError("No se encontró el libro de aberturas.")
    if origen.stat().st_size > TAMANO_MAXIMO_BYTES:
        raise ValueError("El libro de aberturas supera 10 MB.")
    try:
        libro = load_workbook(origen, data_only=False, read_only=True)
    except (OSError, BadZipFile, InvalidFileException, ValueError) as error:
        raise ValueError("No se pudo abrir el libro de aberturas.") from error
    try:
        hoja = libro.worksheets[0]
        filas = list(hoja.iter_rows(values_only=True))
    finally:
        libro.close()
    if not filas:
        raise ValueError("El libro de aberturas está vacío.")
    _validar_encabezado(filas[0])
    tamices: list[tuple[float, float]] = []
    fondo = 0.0
    hubo_fondo = False
    for numero, fila in enumerate(filas[1:], start=2):
        abertura, porcentaje = _celdas(fila)
        if abertura is None and porcentaje is None:
            continue
        if hubo_fondo:
            raise ValueError(f"Fila {numero}: el fondo debe ser la última fila con datos.")
        if isinstance(abertura, str) and _normalizar(abertura) == "fondo":
            if porcentaje is None:
                raise ValueError(f"Fila {numero}: el fondo requiere porcentaje retenido.")
            fondo = _numero(porcentaje, numero, "porcentaje")
            hubo_fondo = True
            continue
        if abertura is None or porcentaje is None:
            raise ValueError(f"Fila {numero}: se requieren abertura y porcentaje retenido.")
        tamices.append((_numero(abertura, numero, "abertura"), _numero(porcentaje, numero, "porcentaje")))
    if not tamices:
        raise ValueError("El libro no contiene tamices.")
    return LecturaAberturas(origen.name, tuple(tamices), fondo)


def _validar_encabezado(fila) -> None:
    if fila is None or len(fila) < 2:
        raise ValueError("La primera fila debe titular Abertura (mm) y % Retenido.")
    abertura = _normalizar(fila[0])
    retenido = _normalizar(fila[1])
    if "abertura" not in abertura or "retenido" not in retenido:
        raise ValueError("La primera fila debe titular Abertura (mm) y % Retenido.")


def _celdas(fila):
    if fila is None:
        return None, None
    abertura = fila[0] if len(fila) > 0 else None
    porcentaje = fila[1] if len(fila) > 1 else None
    if _vacio(abertura) and _vacio(porcentaje):
        return None, None
    return None if _vacio(abertura) else abertura, None if _vacio(porcentaje) else porcentaje


def _vacio(valor) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def _normalizar(valor) -> str:
    texto = " ".join(str(valor).strip().lower().split())
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def _numero(valor, fila: int, nombre: str) -> float:
    if isinstance(valor, str) and valor.lstrip().startswith("="):
        raise ValueError(f"Fila {fila}: no se admiten fórmulas; guarde valores.")
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        raise ValueError(f"Fila {fila}: el {nombre} debe ser numérico.")
    numero = float(valor)
    if not math.isfinite(numero):
        raise ValueError(f"Fila {fila}: el {nombre} debe ser finito.")
    return numero
